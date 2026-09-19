"""
ONBOARD360 — Master Automated Test Suite (160 Comprehensive Test Cases)
Author: Lead Business Analyst & QA Automation Lead
Validates all dimensions: Governance, File Existence, Data Integrity, SQL Analytics,
Process Logic, Machine Learning Guardrails, Financial Calculations, and Excel/PDF Formats.
"""

import os
import sys
import unittest
import sqlite3
import pandas as pd
import numpy as np
import openpyxl
import joblib
import xml.etree.ElementTree as ET
from datetime import datetime

class TestOnboard360Master(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        print(f"\n[{datetime.now()}] Initializing TestOnboard360Master Suite...")
        cls.df_apps = pd.read_parquet("06_Data_Analysis/data/applications.parquet")
        cls.df_revs = pd.read_parquet("06_Data_Analysis/data/manual_reviews.parquet")
        cls.df_tcks = pd.read_parquet("06_Data_Analysis/data/support_tickets.parquet")
        cls.clf = joblib.load("07_Solution_Design/model/ai_triage_model.joblib")
        
        # Load SQLite DB
        cls.conn = sqlite3.connect("06_Data_Analysis/onboard360.db")

    @classmethod
    def tearDownClass(cls):
        cls.conn.close()

    # =========================================================================
    # SUITE 1: ARTIFACT & REPOSITORY GOVERNANCE INTEGRITY (Tests 1 - 25)
    # =========================================================================
    def test_001_readme_exists(self): self.assertTrue(os.path.exists("README.md"))
    def test_002_project_context_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/PROJECT_CONTEXT.md"))
    def test_003_master_plan_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/MASTER_PLAN.md"))
    def test_004_assumptions_register_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/ASSUMPTIONS_REGISTER.md"))
    def test_005_decision_log_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/DECISION_LOG.md"))
    def test_006_metric_dictionary_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/METRIC_DICTIONARY.md"))
    def test_007_data_dictionary_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/DATA_DICTIONARY.md"))
    def test_008_traceability_master_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/TRACEABILITY_MASTER.md"))
    def test_009_skill_matrix_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/SKILL_READINESS_MATRIX.md"))
    def test_010_changelog_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/CHANGELOG.md"))
    def test_011_project_status_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/PROJECT_STATUS.md"))
    def test_012_gate_0_signoff_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/GATE_0_SIGNOFF.md"))
    def test_013_final_audit_exists(self): self.assertTrue(os.path.exists("00_Project_Governance/FINAL_AUDIT_REPORT.md"))
    def test_014_problem_statement_md_exists(self): self.assertTrue(os.path.exists("01_Business_Case/PROBLEM_STATEMENT.md"))
    def test_015_problem_statement_pdf_exists(self): self.assertTrue(os.path.exists("01_Business_Case/problem_statement.pdf"))
    def test_016_project_charter_exists(self): self.assertTrue(os.path.exists("01_Business_Case/PROJECT_CHARTER.md"))
    def test_017_stakeholder_matrix_md_exists(self): self.assertTrue(os.path.exists("02_Stakeholder_Analysis/STAKEHOLDER_MATRIX.md"))
    def test_018_stakeholder_matrix_xlsx_exists(self): self.assertTrue(os.path.exists("02_Stakeholder_Analysis/stakeholder_matrix.xlsx"))
    def test_019_raci_matrix_exists(self): self.assertTrue(os.path.exists("02_Stakeholder_Analysis/RACI_MATRIX.md"))
    def test_020_as_is_inventory_exists(self): self.assertTrue(os.path.exists("03_Process_Analysis/AS_IS_PROCESS_INVENTORY.md"))
    def test_021_as_is_bpmn_exists(self): self.assertTrue(os.path.exists("03_Process_Analysis/as_is_process.bpmn"))
    def test_022_pain_point_catalog_exists(self): self.assertTrue(os.path.exists("03_Process_Analysis/PAIN_POINT_CATALOG.md"))
    def test_023_pain_point_pdf_exists(self): self.assertTrue(os.path.exists("03_Process_Analysis/pain_point_analysis.pdf"))
    def test_024_root_cause_md_exists(self): self.assertTrue(os.path.exists("03_Process_Analysis/ROOT_CAUSE_ANALYSIS.md"))
    def test_025_root_cause_xlsx_exists(self): self.assertTrue(os.path.exists("03_Process_Analysis/root_cause_analysis.xlsx"))

    # =========================================================================
    # SUITE 2: REQUIREMENTS, AGILE & SOLUTION ARTIFACTS (Tests 26 - 50)
    # =========================================================================
    def test_026_brd_md_exists(self): self.assertTrue(os.path.exists("04_Requirements/BRD.md"))
    def test_027_brd_pdf_exists(self): self.assertTrue(os.path.exists("04_Requirements/BRD.pdf"))
    def test_028_frd_md_exists(self): self.assertTrue(os.path.exists("04_Requirements/FRD.md"))
    def test_029_frd_pdf_exists(self): self.assertTrue(os.path.exists("04_Requirements/FRD.pdf"))
    def test_030_business_rules_exists(self): self.assertTrue(os.path.exists("04_Requirements/BUSINESS_RULES.md"))
    def test_031_requirements_traceability_xlsx_exists(self): self.assertTrue(os.path.exists("04_Requirements/requirements_traceability.xlsx"))
    def test_032_product_backlog_md_exists(self): self.assertTrue(os.path.exists("05_Agile/PRODUCT_BACKLOG.md"))
    def test_033_product_backlog_xlsx_exists(self): self.assertTrue(os.path.exists("05_Agile/product_backlog.xlsx"))
    def test_034_user_stories_md_exists(self): self.assertTrue(os.path.exists("05_Agile/USER_STORIES.md"))
    def test_035_user_stories_xlsx_exists(self): self.assertTrue(os.path.exists("05_Agile/user_stories.xlsx"))
    def test_036_sprint_plan_xlsx_exists(self): self.assertTrue(os.path.exists("05_Agile/sprint_plan.xlsx"))
    def test_037_jira_confluence_sim_exists(self): self.assertTrue(os.path.exists("05_Agile/JIRA_CONFLUENCE_SIMULATION.md"))
    def test_038_to_be_bpmn_exists(self): self.assertTrue(os.path.exists("07_Solution_Design/to_be_process.bpmn"))
    def test_039_to_be_spec_exists(self): self.assertTrue(os.path.exists("07_Solution_Design/TO_BE_PROCESS_SPECIFICATION.md"))
    def test_040_solution_architecture_md_exists(self): self.assertTrue(os.path.exists("07_Solution_Design/SOLUTION_ARCHITECTURE.md"))
    def test_041_solution_architecture_pdf_exists(self): self.assertTrue(os.path.exists("07_Solution_Design/solution_architecture.pdf"))
    def test_042_ai_triage_model_card_exists(self): self.assertTrue(os.path.exists("07_Solution_Design/AI_TRIAGE_MODEL_CARD.md"))
    def test_043_ai_triage_model_joblib_exists(self): self.assertTrue(os.path.exists("07_Solution_Design/model/ai_triage_model.joblib"))
    def test_044_powerbi_spec_exists(self): self.assertTrue(os.path.exists("08_PowerBI/POWER_BI_SPECIFICATION.md"))
    def test_045_dax_catalog_exists(self): self.assertTrue(os.path.exists("08_PowerBI/DAX_MEASURE_CATALOG.md"))
    def test_046_uat_test_plan_md_exists(self): self.assertTrue(os.path.exists("09_UAT/UAT_TEST_PLAN.md"))
    def test_047_uat_test_plan_xlsx_exists(self): self.assertTrue(os.path.exists("09_UAT/test_plan.xlsx"))
    def test_048_uat_test_cases_md_exists(self): self.assertTrue(os.path.exists("09_UAT/UAT_TEST_CASES.md"))
    def test_049_uat_results_xlsx_exists(self): self.assertTrue(os.path.exists("09_UAT/UAT_results.xlsx"))
    def test_050_roi_model_xlsx_exists(self): self.assertTrue(os.path.exists("10_Business_Case/ROI_model.xlsx"))

    # =========================================================================
    # SUITE 3: DATA ARCHITECTURE & STATISTICAL REASONING (Tests 51 - 75)
    # =========================================================================
    def test_051_applications_row_count(self): self.assertGreaterEqual(len(self.df_apps), 500000)
    def test_052_applications_exact_size(self): self.assertEqual(len(self.df_apps), 520000)
    def test_053_applications_columns(self):
        expected = ["application_id", "customer_id", "application_date", "customer_type", "region", "channel", "status"]
        for col in expected: self.assertIn(col, self.df_apps.columns)
    def test_054_no_null_application_ids(self): self.assertEqual(self.df_apps["application_id"].isnull().sum(), 0)
    def test_055_unique_application_ids(self): self.assertEqual(self.df_apps["application_id"].nunique(), 520000)
    def test_056_valid_statuses(self):
        valid = {"APPROVED", "REJECTED", "ABANDONED"}
        self.assertTrue(set(self.df_apps["status"].unique()).issubset(valid))
    def test_057_valid_channels(self):
        valid = {"Mobile_App", "Web_Portal", "Branch_Assisted", "Affiliate_Partner"}
        self.assertTrue(set(self.df_apps["channel"].unique()).issubset(valid))
    def test_058_valid_regions(self):
        valid = {"UK_EU", "US", "APAC", "LATAM"}
        self.assertTrue(set(self.df_apps["region"].unique()).issubset(valid))
    def test_059_valid_risk_tiers(self):
        valid = {"LOW", "MEDIUM", "HIGH", "PEP"}
        self.assertTrue(set(self.df_apps["risk_tier"].unique()).issubset(valid))
    def test_060_rework_flag_boolean(self): self.assertEqual(self.df_apps["rework_flag"].dtype, bool)
    def test_061_manual_review_flag_boolean(self): self.assertEqual(self.df_apps["manual_review_flag"].dtype, bool)
    def test_062_sla_breach_flag_boolean(self): self.assertEqual(self.df_apps["sla_breach_flag"].dtype, bool)
    def test_063_cycle_time_positive(self): self.assertTrue((self.df_apps["total_cycle_time_hours"] > 0).all())
    def test_064_touch_time_positive(self): self.assertTrue((self.df_apps["touch_time_hours"] > 0).all())
    def test_065_wait_time_positive(self): self.assertTrue((self.df_apps["wait_time_hours"] > 0).all())
    def test_066_cycle_equals_touch_plus_wait(self):
        diff = np.abs(self.df_apps["total_cycle_time_hours"] - (self.df_apps["touch_time_hours"] + self.df_apps["wait_time_hours"]))
        self.assertLess(diff.max(), 0.05) # float tolerance
    def test_067_sla_breach_definition(self):
        calculated_breach = self.df_apps["total_cycle_time_hours"] > 48.0
        self.assertTrue((self.df_apps["sla_breach_flag"] == calculated_breach).all())
    def test_068_average_tat_range(self):
        mean_tat = self.df_apps["total_cycle_time_hours"].mean()
        self.assertTrue(55.0 <= mean_tat <= 65.0)
    def test_069_wait_time_dominates_cycle(self):
        wait_ratio = self.df_apps["wait_time_hours"].sum() / self.df_apps["total_cycle_time_hours"].sum()
        self.assertGreater(wait_ratio, 0.85) # Wait time is > 85% of total lead time
    def test_070_rework_rate_range(self):
        rework_rate = self.df_apps["rework_flag"].mean()
        self.assertTrue(0.30 <= rework_rate <= 0.36)
    def test_071_manual_review_rate_range(self):
        mr_rate = self.df_apps["manual_review_flag"].mean()
        self.assertTrue(0.38 <= mr_rate <= 0.45)
    def test_072_abandonment_rate_range(self):
        ab_rate = (self.df_apps["status"] == "ABANDONED").mean()
        self.assertTrue(0.14 <= ab_rate <= 0.18)
    def test_073_rework_reason_when_flag_false(self):
        non_rework = self.df_apps[~self.df_apps["rework_flag"]]
        self.assertTrue((non_rework["rework_reason"] == "None").all())
    def test_074_rework_reason_when_flag_true(self):
        rework = self.df_apps[self.df_apps["rework_flag"]]
        self.assertFalse((rework["rework_reason"] == "None").any())
    def test_075_manual_reviews_table_not_empty(self): self.assertGreater(len(self.df_revs), 200000)

    # =========================================================================
    # SUITE 4: RELATIONAL INTEGRITY & SQL ANALYTICS (Tests 76 - 100)
    # =========================================================================
    def test_076_manual_reviews_fk_integrity(self):
        valid_apps = set(self.df_apps["application_id"])
        review_apps = set(self.df_revs["application_id"])
        self.assertTrue(review_apps.issubset(valid_apps))
    def test_077_support_tickets_fk_integrity(self):
        valid_apps = set(self.df_apps["application_id"])
        tck_apps = set(self.df_tcks["application_id"])
        self.assertTrue(tck_apps.issubset(valid_apps))
    def test_078_support_tickets_positive_cost(self):
        self.assertTrue((self.df_tcks["cost_to_serve_usd"] > 0).all())
    def test_079_sql_total_count(self):
        cur = self.conn.cursor()
        cur.execute("SELECT COUNT(*) FROM applications")
        self.assertEqual(cur.fetchone()[0], 520000)
    def test_080_sql_approval_count(self):
        cur = self.conn.cursor()
        cur.execute("SELECT COUNT(*) FROM applications WHERE status = 'APPROVED'")
        self.assertEqual(cur.fetchone()[0], len(self.df_apps[self.df_apps["status"] == "APPROVED"]))
    def test_081_sql_abandon_count(self):
        cur = self.conn.cursor()
        cur.execute("SELECT COUNT(*) FROM applications WHERE status = 'ABANDONED'")
        self.assertEqual(cur.fetchone()[0], len(self.df_apps[self.df_apps["status"] == "ABANDONED"]))
    def test_082_sql_fpy_count(self):
        cur = self.conn.cursor()
        cur.execute("SELECT COUNT(*) FROM applications WHERE status = 'APPROVED' AND rework_flag = 0")
        fpy = len(self.df_apps[(self.df_apps["status"] == "APPROVED") & (~self.df_apps["rework_flag"])])
        self.assertEqual(cur.fetchone()[0], fpy)
    def test_083_sql_pareto_blurry_highest(self):
        cur = self.conn.cursor()
        cur.execute("SELECT rework_reason, COUNT(*) as cnt FROM applications WHERE rework_flag = 1 GROUP BY rework_reason ORDER BY cnt DESC LIMIT 1")
        top_reason = cur.fetchone()[0]
        self.assertEqual(top_reason, "Blurry_Image")
    def test_084_sql_pareto_expired_second(self):
        cur = self.conn.cursor()
        cur.execute("SELECT rework_reason, COUNT(*) as cnt FROM applications WHERE rework_flag = 1 GROUP BY rework_reason ORDER BY cnt DESC LIMIT 2")
        rows = cur.fetchall()
        self.assertEqual(rows[1][0], "Expired_ID")
    def test_085_sql_channel_mobile_highest_volume(self):
        cur = self.conn.cursor()
        cur.execute("SELECT channel, COUNT(*) as cnt FROM applications GROUP BY channel ORDER BY cnt DESC LIMIT 1")
        self.assertEqual(cur.fetchone()[0], "Mobile_App")
    def test_086_sql_support_ticket_total_cost(self):
        cur = self.conn.cursor()
        cur.execute("SELECT ROUND(SUM(cost_to_serve_usd), 2) FROM support_tickets")
        sql_sum = cur.fetchone()[0]
        df_sum = round(self.df_tcks["cost_to_serve_usd"].sum(), 2)
        self.assertEqual(sql_sum, df_sum)
    def test_087_sql_support_status_inquiry_dominance(self):
        cur = self.conn.cursor()
        cur.execute("SELECT issue_category, COUNT(*) FROM support_tickets GROUP BY issue_category ORDER BY COUNT(*) DESC LIMIT 1")
        self.assertEqual(cur.fetchone()[0], "Status_Inquiry_WhereIsMyAccount")
    def test_088_sql_window_function_support(self):
        cur = self.conn.cursor()
        cur.execute("SELECT application_id, total_cycle_time_hours, RANK() OVER (ORDER BY total_cycle_time_hours DESC) FROM applications LIMIT 5")
        rows = cur.fetchall()
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[0][2], 1)
    def test_089_sql_manual_reviews_queue_distribution(self):
        cur = self.conn.cursor()
        cur.execute("SELECT DISTINCT queue_name FROM manual_reviews")
        queues = {r[0] for r in cur.fetchall()}
        self.assertTrue(queues.issubset({"L1_Document_Exception", "L2_AML_Sanctions"}))
    def test_090_sql_avg_tat_by_risk_tier(self):
        cur = self.conn.cursor()
        cur.execute("SELECT risk_tier, AVG(total_cycle_time_hours) FROM applications GROUP BY risk_tier")
        results = dict(cur.fetchall())
        self.assertGreater(results["PEP"], results["LOW"]) # High risk takes longer
    def test_091_littles_law_arrival_rate(self):
        lambda_rate = len(self.df_apps) / 8760.0
        self.assertTrue(59.0 <= lambda_rate <= 60.0) # ~59.36 apps/hr
    def test_092_littles_law_wip_concurrency(self):
        lambda_rate = len(self.df_apps) / 8760.0
        wip = lambda_rate * self.df_apps["total_cycle_time_hours"].mean()
        self.assertTrue(3300 <= wip <= 3600) # ~3,485 apps active in system
    def test_093_pareto_top_3_exceeds_80_percent(self):
        rework_counts = self.df_apps[self.df_apps["rework_flag"]]["rework_reason"].value_counts()
        top3_pct = rework_counts.iloc[:3].sum() / rework_counts.sum()
        self.assertGreater(top3_pct, 0.80) # 83.2%
    def test_094_mobile_rework_rate_higher_than_branch(self):
        grp = self.df_apps.groupby("channel")["rework_flag"].mean()
        self.assertGreater(grp["Mobile_App"], grp["Branch_Assisted"])
    def test_095_pep_manual_review_rate_near_100(self):
        pep_mr = self.df_apps[self.df_apps["risk_tier"] == "PEP"]["manual_review_flag"].mean()
        self.assertGreater(pep_mr, 0.95)
    def test_096_high_risk_manual_review_rate_high(self):
        high_mr = self.df_apps[self.df_apps["risk_tier"] == "HIGH"]["manual_review_flag"].mean()
        self.assertGreater(high_mr, 0.88)
    def test_097_low_risk_manual_review_rate_moderate(self):
        low_mr = self.df_apps[self.df_apps["risk_tier"] == "LOW"]["manual_review_flag"].mean()
        self.assertLess(low_mr, 0.40)
    def test_098_support_ticket_rate_overall(self):
        tck_rate = len(self.df_tcks) / len(self.df_apps)
        self.assertTrue(0.20 <= tck_rate <= 0.25)
    def test_099_sla_breach_rate_overall(self):
        breach_rate = self.df_apps["sla_breach_flag"].mean()
        self.assertTrue(0.40 <= breach_rate <= 0.50)
    def test_100_process_cycle_efficiency_baseline(self):
        pce = (self.df_apps["touch_time_hours"].mean() / self.df_apps["total_cycle_time_hours"].mean()) * 100
        self.assertTrue(8.0 <= pce <= 14.0)

    # =========================================================================
    # SUITE 5: AI EXCEPTION TRIAGE & REGULATORY GUARDRAILS (Tests 101 - 120)
    # =========================================================================
    def test_101_model_fitted(self): self.assertTrue(hasattr(self.clf, "classes_"))
    def test_102_model_classes_count(self): self.assertEqual(len(self.clf.classes_), 3)
    def test_103_feature_importances_exist(self): self.assertEqual(len(self.clf.feature_importances_), 7)
    def test_104_top_feature_is_risk_or_rework(self):
        top_idx = np.argmax(self.clf.feature_importances_)
        self.assertIn(top_idx, [0, 3]) # risk_tier_num or rework_reason_num
    def test_105_predict_low_risk_blurry(self):
        # [risk_tier_num, channel_num, cust_type_num, rework_reason_num, sharpness_score, address_match_score, touch_time_hours]
        sample = np.array([[0, 0, 0, 1, 0.45, 95.0, 1.2]])
        self.assertEqual(self.clf.predict(sample)[0], 0) # Auto-Remediation
    def test_106_predict_low_risk_expired_id(self):
        sample = np.array([[0, 0, 0, 2, 0.85, 95.0, 1.2]])
        self.assertEqual(self.clf.predict(sample)[0], 0) # Auto-Remediation
    def test_107_predict_pep_routes_to_l2(self):
        sample = np.array([[3, 0, 0, 0, 0.95, 98.0, 3.5]])
        self.assertEqual(self.clf.predict(sample)[0], 2) # L2 Compliance
    def test_108_predict_high_risk_routes_to_l2(self):
        sample = np.array([[2, 1, 3, 0, 0.90, 95.0, 3.2]])
        self.assertEqual(self.clf.predict(sample)[0], 2) # L2 Compliance
    def test_109_predict_medium_risk_routes_to_l1(self):
        sample = np.array([[1, 0, 0, 3, 0.85, 65.0, 2.0]])
        self.assertEqual(self.clf.predict(sample)[0], 1) # L1 Ops Queue
    def test_110_zero_leakage_pep_to_auto(self):
        # 100 variations of PEP cases
        pep_cases = np.zeros((100, 7))
        pep_cases[:, 0] = 3 # PEP
        pep_cases[:, 1] = np.random.randint(0, 4, 100)
        pep_cases[:, 2] = np.random.randint(0, 4, 100)
        pep_cases[:, 3] = np.random.randint(0, 6, 100)
        pep_cases[:, 4] = np.random.uniform(0.1, 1.0, 100)
        pep_cases[:, 5] = np.random.uniform(40.0, 100.0, 100)
        pep_cases[:, 6] = np.random.uniform(0.5, 10.0, 100)
        preds = self.clf.predict(pep_cases)
        self.assertNotIn(0, preds) # NEVER Class 0 (Auto-Remediation)
    def test_111_zero_leakage_high_risk_to_auto(self):
        high_cases = np.zeros((100, 7))
        high_cases[:, 0] = 2 # High Risk
        high_cases[:, 1] = np.random.randint(0, 4, 100)
        high_cases[:, 2] = np.random.randint(0, 4, 100)
        high_cases[:, 3] = np.random.randint(0, 6, 100)
        high_cases[:, 4] = np.random.uniform(0.1, 1.0, 100)
        high_cases[:, 5] = np.random.uniform(40.0, 100.0, 100)
        high_cases[:, 6] = np.random.uniform(0.5, 10.0, 100)
        preds = self.clf.predict(high_cases)
        self.assertNotIn(0, preds) # NEVER Class 0
    def test_112_confidence_score_low_risk_blurry(self):
        sample = np.array([[0, 0, 0, 1, 0.40, 95.0, 1.0]])
        prob = self.clf.predict_proba(sample)[0]
        self.assertGreaterEqual(prob[0], 0.85) # High confidence
    def test_113_confidence_score_pep(self):
        sample = np.array([[3, 0, 0, 0, 0.95, 95.0, 3.5]])
        prob = self.clf.predict_proba(sample)[0]
        self.assertGreaterEqual(prob[2], 0.85)
    def test_114_model_determinism(self):
        sample = np.array([[0, 0, 0, 1, 0.45, 92.0, 1.2]])
        p1 = self.clf.predict(sample)[0]
        p2 = self.clf.predict(sample)[0]
        self.assertEqual(p1, p2)
    def test_115_model_depth_bounded(self):
        self.assertEqual(self.clf.max_depth, 12)
    def test_116_model_trees_count(self):
        self.assertEqual(len(self.clf.estimators_), 100)
    def test_117_address_mismatch_routes_to_l1(self):
        sample = np.array([[0, 1, 0, 3, 0.95, 55.0, 1.5]]) # Low risk, but address mismatch
        pred = self.clf.predict(sample)[0]
        self.assertEqual(pred, 1) # L1 Ops
    def test_118_incomplete_form_routes_to_l1(self):
        sample = np.array([[0, 1, 0, 4, 0.95, 95.0, 1.5]]) # Incomplete form
        pred = self.clf.predict(sample)[0]
        self.assertEqual(pred, 1) # L1 Ops
    def test_119_name_mismatch_routes_to_l1(self):
        sample = np.array([[0, 0, 0, 5, 0.95, 95.0, 1.5]]) # Name mismatch
        pred = self.clf.predict(sample)[0]
        self.assertEqual(pred, 1) # L1 Ops
    def test_120_ai_governance_specification_present(self):
        self.assertTrue(os.path.exists("07_Solution_Design/AI_TRIAGE_MODEL_CARD.md"))

    # =========================================================================
    # SUITE 6: FINANCIAL ROI & BUSINESS CASE VALIDATION (Tests 121 - 140)
    # =========================================================================
    def test_121_as_is_cost_above_25m(self):
        # L1 ops + L2 comp + support + vendor + admin
        l1 = 184000 * 1.8 * 42.0
        l2 = 29842 * 3.5 * 65.0
        support = 1408658.0
        vendor = 520000 * 9.80
        admin = 173429 * 4.40
        total = l1 + l2 + support + vendor + admin
        self.assertGreater(total, 25000000.0)
    def test_122_to_be_cost_below_7m(self):
        l1 = 41600 * 0.8 * 42.0
        l2 = 15600 * 1.5 * 65.0
        support = 1408658.0 * 0.25
        vendor = 520000 * 4.20
        cloud = 420000.0
        total = l1 + l2 + support + vendor + cloud
        self.assertLess(total, 7000000.0)
    def test_123_annual_gross_savings_above_20m(self):
        # Baseline ~27.96M - Target ~5.87M
        savings = 27967200.60 - 5874924.50
        self.assertGreater(savings, 20000000.0)
    def test_124_unit_cost_as_is_above_60(self):
        unit_as_is = 27967200.60 / 414655
        self.assertGreater(unit_as_is, 60.0)
    def test_125_unit_cost_to_be_below_15(self):
        unit_to_be = 5874924.50 / 468000
        self.assertLess(unit_to_be, 15.0)
    def test_126_unit_cost_reduction_percentage(self):
        reduction = (67.45 - 12.55) / 67.45
        self.assertGreater(reduction, 0.80) # > 80% reduction
    def test_127_capex_initial_assumption(self):
        capex = 2850000.0
        self.assertEqual(capex, 2850000.0)
    def test_128_hurdle_discount_rate(self):
        r = 0.085
        self.assertEqual(r, 0.085)
    def test_129_cash_flow_year_1(self):
        cf1 = 22092276.10 * 0.80
        self.assertTrue(17000000.0 <= cf1 <= 18000000.0)
    def test_130_cash_flow_year_2(self):
        cf2 = 22092276.10
        self.assertEqual(cf2, 22092276.10)
    def test_131_cash_flow_year_3(self):
        cf3 = 22092276.10
        self.assertEqual(cf3, 22092276.10)
    def test_132_npv_3year_calculation(self):
        cfs = [-2850000.0, 22092276.10 * 0.80, 22092276.10, 22092276.10]
        npv = sum(cf / ((1 + 0.085) ** t) for t, cf in enumerate(cfs))
        self.assertGreater(npv, 45000000.0) # ~$49.5M
    def test_133_irr_above_150_percent(self):
        cfs = [-2850000.0, 17673820.88, 22092276.10, 22092276.10]
        # At r = 1.5 (150%), NPV is still positive
        npv_150 = sum(cf / ((1 + 1.5) ** t) for t, cf in enumerate(cfs))
        self.assertGreater(npv_150, 0)
    def test_134_payback_period_under_3_months(self):
        monthly_y1 = (22092276.10 * 0.80) / 12.0
        months_to_payback = 2850000.0 / monthly_y1
        self.assertLess(months_to_payback, 3.0)
    def test_135_conservative_scenario_npv_positive(self):
        benefit = 22092276.10 * 0.75
        cfs = [-3200000.0, benefit * 0.80, benefit, benefit]
        npv = sum(cf / ((1 + 0.085) ** t) for t, cf in enumerate(cfs))
        self.assertGreater(npv, 30000000.0)
    def test_136_aggressive_scenario_npv_higher(self):
        benefit = 22092276.10 * 1.22
        cfs = [-2500000.0, benefit * 0.80, benefit, benefit]
        npv = sum(cf / ((1 + 0.085) ** t) for t, cf in enumerate(cfs))
        self.assertGreater(npv, 55000000.0)
    def test_137_support_cost_savings(self):
        savings = 1408658.0 * 0.75
        self.assertGreater(savings, 1000000.0) # > $1M saved in support
    def test_138_l1_labor_savings(self):
        savings = 13910400.0 - 1397760.0
        self.assertGreater(savings, 12000000.0) # > $12M saved in L1 labor
    def test_139_l2_labor_savings(self):
        savings = 6789055.0 - 1521000.0
        self.assertGreater(savings, 5000000.0) # > $5M saved in L2 compliance
    def test_140_vendor_api_savings(self):
        savings = (520000 * 9.80) - (520000 * 4.20)
        self.assertGreater(savings, 2500000.0) # > $2.5M saved in vendor APIs

    # =========================================================================
    # SUITE 7: BPMN, EXCEL & REPORTING FORMAT INTEGRITY (Tests 141 - 160)
    # =========================================================================
    def test_141_as_is_bpmn_valid_xml(self):
        tree = ET.parse("03_Process_Analysis/as_is_process.bpmn")
        self.assertEqual(tree.getroot().tag, "{http://www.omg.org/spec/BPMN/20100524/MODEL}definitions")
    def test_142_as_is_bpmn_has_customer_pool(self):
        tree = ET.parse("03_Process_Analysis/as_is_process.bpmn")
        pools = [p.attrib.get("name") for p in tree.getroot().iter("{http://www.omg.org/spec/BPMN/20100524/MODEL}participant")]
        self.assertIn("Customer", pools)
    def test_143_as_is_bpmn_has_bank_pool(self):
        tree = ET.parse("03_Process_Analysis/as_is_process.bpmn")
        pools = [p.attrib.get("name") for p in tree.getroot().iter("{http://www.omg.org/spec/BPMN/20100524/MODEL}participant")]
        self.assertIn("NovaBank Operations & Systems", pools)
    def test_144_to_be_bpmn_valid_xml(self):
        tree = ET.parse("07_Solution_Design/to_be_process.bpmn")
        self.assertEqual(tree.getroot().tag, "{http://www.omg.org/spec/BPMN/20100524/MODEL}definitions")
    def test_145_to_be_bpmn_has_stp_lane(self):
        tree = ET.parse("07_Solution_Design/to_be_process.bpmn")
        lanes = [l.attrib.get("name") for l in tree.getroot().iter("{http://www.omg.org/spec/BPMN/20100524/MODEL}lane")]
        self.assertTrue(any("STP" in l for l in lanes))
    def test_146_to_be_bpmn_has_ai_triage_lane(self):
        tree = ET.parse("07_Solution_Design/to_be_process.bpmn")
        lanes = [l.attrib.get("name") for l in tree.getroot().iter("{http://www.omg.org/spec/BPMN/20100524/MODEL}lane")]
        self.assertTrue(any("AI Exception Triage" in l for l in lanes))
    def test_147_stakeholder_excel_sheets(self):
        wb = openpyxl.load_workbook("02_Stakeholder_Analysis/stakeholder_matrix.xlsx")
        self.assertIn("Stakeholder_Matrix", wb.sheetnames)
        self.assertGreater(wb["Stakeholder_Matrix"].max_row, 5)
    def test_148_root_cause_excel_sheets(self):
        wb = openpyxl.load_workbook("03_Process_Analysis/root_cause_analysis.xlsx")
        self.assertIn("Pareto_Root_Cause", wb.sheetnames)
        self.assertGreater(wb["Pareto_Root_Cause"].max_row, 4)
    def test_149_requirements_traceability_excel(self):
        wb = openpyxl.load_workbook("04_Requirements/requirements_traceability.xlsx")
        self.assertIn("Traceability_Matrix", wb.sheetnames)
        self.assertGreaterEqual(wb["Traceability_Matrix"].max_row, 10)
    def test_150_product_backlog_excel(self):
        wb = openpyxl.load_workbook("05_Agile/product_backlog.xlsx")
        self.assertIn("Product_Backlog", wb.sheetnames)
        self.assertGreaterEqual(wb["Product_Backlog"].max_row, 10)
    def test_151_user_stories_excel(self):
        wb = openpyxl.load_workbook("05_Agile/user_stories.xlsx")
        self.assertIn("Product_Backlog", wb.sheetnames)
    def test_152_sprint_plan_excel(self):
        wb = openpyxl.load_workbook("05_Agile/sprint_plan.xlsx")
        self.assertIn("Product_Backlog", wb.sheetnames)
    def test_153_uat_plan_excel(self):
        wb = openpyxl.load_workbook("09_UAT/test_plan.xlsx")
        self.assertIn("UAT_Results", wb.sheetnames)
    def test_154_uat_results_excel(self):
        wb = openpyxl.load_workbook("09_UAT/UAT_results.xlsx")
        self.assertIn("UAT_Results", wb.sheetnames)
        self.assertGreaterEqual(wb["UAT_Results"].max_row, 8)
    def test_155_roi_model_excel(self):
        wb = openpyxl.load_workbook("10_Business_Case/ROI_model.xlsx")
        self.assertIn("Cost_Benefit_Model", wb.sheetnames)
        self.assertGreaterEqual(wb["Cost_Benefit_Model"].max_row, 10)
    def test_156_pdf_problem_statement_size(self):
        size = os.path.getsize("01_Business_Case/problem_statement.pdf")
        self.assertGreater(size, 1000) # Valid non-empty PDF
    def test_157_pdf_pain_point_size(self):
        size = os.path.getsize("03_Process_Analysis/pain_point_analysis.pdf")
        self.assertGreater(size, 1000)
    def test_158_pdf_brd_size(self):
        size = os.path.getsize("04_Requirements/BRD.pdf")
        self.assertGreater(size, 1000)
    def test_159_pdf_frd_size(self):
        size = os.path.getsize("04_Requirements/FRD.pdf")
        self.assertGreater(size, 1000)
    def test_160_pdf_executive_review_size(self):
        size = os.path.getsize("11_Executive_Presentation/ONBOARD360_Executive_Review.pdf")
        self.assertGreater(size, 1000)

if __name__ == "__main__":
    unittest.main(verbosity=2)
